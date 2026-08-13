import sys, zipfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(r"C:\Users\affla\OneDrive\Documents\Saas\Claude projects\COCOPALM")
USER = ROOT / "Coco-Palms-Export-Profile.pptx"
MINE = ROOT / "build" / "Coco-Palms-Export-Profile.pptx"


def media(p):
    with zipfile.ZipFile(p) as z:
        return [(n, z.getinfo(n).file_size)
                for n in z.namelist() if n.startswith("ppt/media/")]


def survey(p, tag):
    pr = Presentation(p)
    print("=" * 70)
    print("%s  |  %d slides  |  %.2f MB" % (tag, len(pr.slides), p.stat().st_size / 1024 / 1024))
    m = media(p)
    print("media parts: %d, total %.2f MB" % (len(m), sum(s for _, s in m) / 1024 / 1024))
    return pr


def texts(pr):
    """slide index -> list of text strings (shapes + tables)"""
    out = {}
    for i, s in enumerate(pr.slides, 1):
        vals = []
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                vals.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        if c.text.strip():
                            vals.append("[cell] " + c.text.strip())
        out[i] = vals
    return out


def counts(pr):
    out = {}
    for i, s in enumerate(pr.slides, 1):
        pics = sum(1 for sh in s.shapes if sh.shape_type == 13)
        out[i] = (len(s.shapes), pics)
    return out


pu = survey(USER, "YOUR VERSION (root)")
pm = survey(MINE, "MY BUILD (build/)")

tu, tm = texts(pu), texts(pm)
cu, cm = counts(pu), counts(pm)

print()
print("=" * 70)
print("PER-SLIDE COMPARISON   (shapes / pictures)")
print("=" * 70)
n = max(len(tu), len(tm))
for i in range(1, n + 1):
    su = cu.get(i, (0, 0))
    sm = cm.get(i, (0, 0))
    flag = "  <-- CHANGED" if su != sm else ""
    print("slide %2d   yours %2d/%d   mine %2d/%d%s" % (i, su[0], su[1], sm[0], sm[1], flag))

print()
print("=" * 70)
print("TEXT ADDED IN YOUR VERSION (not in mine)")
print("=" * 70)
for i in range(1, n + 1):
    mine_set = set(tm.get(i, []))
    added = [t for t in tu.get(i, []) if t not in mine_set]
    if added:
        print("\n--- slide %d ---" % i)
        for t in added:
            print("   + " + t.replace("\n", " | ")[:150])

print()
print("=" * 70)
print("TEXT REMOVED IN YOUR VERSION (was in mine)")
print("=" * 70)
for i in range(1, n + 1):
    yours_set = set(tu.get(i, []))
    gone = [t for t in tm.get(i, []) if t not in yours_set]
    if gone:
        print("\n--- slide %d ---" % i)
        for t in gone:
            print("   - " + t.replace("\n", " | ")[:150])
