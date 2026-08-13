import shutil
from pathlib import Path
from pptx import Presentation

ROOT = Path(r"C:\Users\affla\OneDrive\Documents\Saas\Claude projects\COCOPALM")
USER = ROOT / "Coco-Palms-Export-Profile.pptx"
BAK = ROOT / "build" / "Coco-Palms-Export-Profile.USER-BACKUP.pptx"

# back up their edited file before touching it
if not BAK.exists():
    shutil.copy2(USER, BAK)
    print("backup written -> build/%s" % BAK.name)

pr = Presentation(USER)
fixed = []

for i, s in enumerate(pr.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                # assign run.text (not text_frame.text) so the run's formatting survives
                if "BELEIVE" in r.text:
                    before = r.text
                    r.text = r.text.replace("BELEIVE", "BELIEVE")
                    fixed.append((i, before, r.text))
                elif "Beleive" in r.text:
                    before = r.text
                    r.text = r.text.replace("Beleive", "Believe")
                    fixed.append((i, before, r.text))

if fixed:
    pr.save(USER)
    for i, b, a in fixed:
        print("slide %d: %r -> %r" % (i, b, a))
    print("saved.")
else:
    print("no spelling fix applied - string not found")
