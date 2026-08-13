import zipfile
from pathlib import Path
from pptx import Presentation

ROOT = Path(r"C:\Users\affla\OneDrive\Documents\Saas\Claude projects\COCOPALM")
USER = ROOT / "Coco-Palms-Export-Profile.pptx"

pr = Presentation(USER)

print("=== FULL TEXT, SLIDE 2 ===")
for sh in pr.slides[1].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print("  * " + sh.text_frame.text.strip().replace("\n", " | "))

print()
print("=== FULL TEXT, SLIDE 12 ===")
for sh in pr.slides[11].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print("  * " + sh.text_frame.text.strip().replace("\n", " | "))

print()
print("=== MEDIA PARTS (yours) ===")
with zipfile.ZipFile(USER) as z:
    for n in sorted(z.namelist()):
        if n.startswith("ppt/media/"):
            print("  %-34s %8.1f KB" % (n.replace("ppt/media/", ""), z.getinfo(n).file_size / 1024))

print()
print("=== which slide uses which media ===")
for i, s in enumerate(pr.slides, 1):
    for sh in s.shapes:
        if sh.shape_type == 13:  # PICTURE
            try:
                nm = sh.image.blob
                print("  slide %2d  %-28s %7.1f KB  %sx%s"
                      % (i, Path(sh.image.filename or "embedded").name,
                         len(nm) / 1024, sh.image.size[0], sh.image.size[1]))
            except Exception as e:
                print("  slide %2d  <unreadable image: %s>" % (i, e))
