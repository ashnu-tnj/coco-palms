import sys, os
from pptx import Presentation
from pptx.util import Emu

DECK = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Coco-Palms-Export-Profile.pptx")
p = Presentation(DECK)
SW = Emu(p.slide_width).inches
SH = Emu(p.slide_height).inches
print("OPENS OK.  slides:", len(p.slides), " size: %.3f x %.3f in" % (SW, SH))

problems = []
for i, s in enumerate(p.slides, 1):
    print("\n=== SLIDE %d  (%d shapes) ===" % (i, len(s.shapes)))
    for sh in s.shapes:
        try:
            x, y = Emu(sh.left).inches, Emu(sh.top).inches
            w, h = Emu(sh.width).inches, Emu(sh.height).inches
        except Exception:
            continue
        # off-canvas / out-of-bounds geometry check
        if x < -0.01 or y < -0.01 or x + w > SW + 0.01 or y + h > SH + 0.01:
            problems.append("S%d %s out of bounds: x=%.2f y=%.2f w=%.2f h=%.2f"
                            % (i, sh.shape_type, x, y, w, h))
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                print("   [%.2f,%.2f %.2fx%.2f] %s" % (x, y, w, h, t.replace("\n", " | ")[:130]))

print("\n\n===== GEOMETRY PROBLEMS =====")
if problems:
    for q in problems:
        print(" !", q)
else:
    print(" none - every shape sits inside the slide")

# placeholder-leak check (the AI-slop patterns, not our deliberate TO CONFIRM)
import re
bad = re.compile(r"\bx{3,}\b|lorem|ipsum|\bTODO\b|\[insert|placeholder text", re.I)
leaks = []
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and bad.search(sh.text_frame.text):
            leaks.append((i, sh.text_frame.text[:80]))
print("\n===== PLACEHOLDER LEAKS =====")
print(leaks if leaks else " none")

tbd = sum(sh.text_frame.text.count("TO CONFIRM")
          for s in p.slides for sh in s.shapes if sh.has_text_frame)
print("\ndeliberate 'TO CONFIRM' slots:", tbd)
